from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet

from . import __version__
from .analyzer import classify_cell
from .models import ConsolidatedGroup, ExtractedPage, ExtractedTable, ExtractionResult


# ---------------------------------------------------------------------------
# Charte graphique
#
# Palette resserree autour d'un bleu nuit (identite / structure) et d'un
# bleu accent (interactions, mise en avant), plus des teintes douces pour
# les statuts (succes / a verifier) - pensee pour rester lisible et sobre
# a l'impression comme a l'ecran, sans neons ni surcharge de couleurs.
# ---------------------------------------------------------------------------
NAVY = "1B2A4A"
ACCENT = "3B6EA5"
ACCENT_LIGHT = "E8F0F9"
ACCENT_SOFT = "D3E3F3"
CARD_BORDER = "B9CFE6"
GRAY_BORDER = "D8DCE3"
GRAY_TEXT = "6B7280"
INK = "1F2933"
WHITE = "FFFFFF"
SUCCESS_BG = "E5F4EA"
SUCCESS_TEXT = "1E7B34"
SUCCESS_BORDER = "BEE3C8"
WARNING_BG = "FDF1DC"
WARNING_TEXT = "B4690E"
WARNING_BORDER = "F2D9AE"

FONT_NAME = "Calibri"

TITLE_FONT = Font(name=FONT_NAME, bold=True, size=20, color=WHITE)
SUBTITLE_FONT = Font(name=FONT_NAME, size=10, color=ACCENT_SOFT)
SECTION_FONT = Font(name=FONT_NAME, bold=True, size=12, color=NAVY)
PAGE_TITLE_FONT = Font(name=FONT_NAME, bold=True, size=15, color=WHITE)
PAGE_SUBTITLE_FONT = Font(name=FONT_NAME, size=9.5, color=INK)
KPI_NUMBER_FONT = Font(name=FONT_NAME, bold=True, size=22, color=NAVY)
KPI_LABEL_FONT = Font(name=FONT_NAME, bold=True, size=8.5, color=GRAY_TEXT)
HEADER_FONT = Font(name=FONT_NAME, bold=True, size=10, color=WHITE)
BODY_FONT = Font(name=FONT_NAME, size=10, color=INK)
BOLD_FONT = Font(name=FONT_NAME, bold=True, size=10, color=INK)
CAPTION_FONT = Font(name=FONT_NAME, italic=True, size=8.5, color=GRAY_TEXT)
TABLE_TITLE_FONT = Font(name=FONT_NAME, bold=True, size=12, color=ACCENT)
HYPERLINK_FONT = Font(name=FONT_NAME, bold=True, size=10, color=ACCENT, underline=None)
BACK_LINK_FONT = Font(name=FONT_NAME, bold=True, size=9, color=ACCENT)
WARNING_FONT = Font(name=FONT_NAME, size=9.5, color=WARNING_TEXT)

NAVY_FILL = PatternFill("solid", fgColor=NAVY)
ACCENT_FILL = PatternFill("solid", fgColor=ACCENT)
ACCENT_LIGHT_FILL = PatternFill("solid", fgColor=ACCENT_LIGHT)
ACCENT_SOFT_FILL = PatternFill("solid", fgColor=ACCENT_SOFT)
WHITE_FILL = PatternFill("solid", fgColor=WHITE)
SUCCESS_FILL = PatternFill("solid", fgColor=SUCCESS_BG)
WARNING_FILL = PatternFill("solid", fgColor=WARNING_BG)

THIN_GRAY = Side(style="thin", color=GRAY_BORDER)
BOX_BORDER = Border(left=THIN_GRAY, right=THIN_GRAY, top=THIN_GRAY, bottom=THIN_GRAY)
THIN_CARD = Side(style="thin", color=CARD_BORDER)
CARD_BORDER_STYLE = Border(left=THIN_CARD, right=THIN_CARD, top=THIN_CARD, bottom=THIN_CARD)
THIN_SUCCESS = Side(style="thin", color=SUCCESS_BORDER)
SUCCESS_BORDER_STYLE = Border(left=THIN_SUCCESS, right=THIN_SUCCESS, top=THIN_SUCCESS, bottom=THIN_SUCCESS)
THIN_WARNING = Side(style="thin", color=WARNING_BORDER)
WARNING_BORDER_STYLE = Border(left=THIN_WARNING, right=THIN_WARNING, top=THIN_WARNING, bottom=THIN_WARNING)
ACCENT_BOTTOM_BORDER = Border(bottom=Side(style="thin", color=ACCENT))

LOW_CONFIDENCE_THRESHOLD = 0.8
TEXT_ZONE_SPAN = 6


def export_to_excel(result: ExtractionResult, output_dir: Path, mode: str = "both") -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Synthese"

    page_sheet_names: dict[int, str] = {}
    for page in result.pages:
        page_sheet_names[page.page_number] = safe_sheet_title(f"Page_{page.page_number:03d}")

    write_summary_sheet(summary_sheet, result, page_sheet_names)
    summary_sheet.sheet_properties.tabColor = NAVY

    used_table_names: set[str] = set()

    if result.consolidated_groups:
        consolidated_sheet = workbook.create_sheet(title="Tableaux_consolides")
        write_consolidated_sheet(consolidated_sheet, result, page_sheet_names, used_table_names)
        consolidated_sheet.sheet_properties.tabColor = NAVY

    for page in result.pages:
        sheet = workbook.create_sheet(title=page_sheet_names[page.page_number])
        write_page_sheet(sheet, page, mode=mode, used_table_names=used_table_names)
        sheet.sheet_properties.tabColor = WARNING_TEXT if page.warnings else ACCENT

    write_metadata_sheet(workbook, result)
    write_quality_sheet(workbook, result, page_sheet_names)

    output_path = output_dir / f"{result.source_path.stem}_extrait.xlsx"
    return save_workbook_safely(workbook, output_path)


# ---------------------------------------------------------------------------
# Cartes indicateurs (KPI) - composant reutilise sur la Synthese et sur la
# feuille Tableaux_consolides, pour garder une charte visuelle identique
# partout ou l'on affiche des indicateurs chiffres.
# ---------------------------------------------------------------------------
def write_kpi_cards(
    sheet: Worksheet,
    kpis: list[tuple[str, object, bool]],
    row_number: int,
    row_label: int,
    start_col: int = 1,
    card_width: int = 2,
) -> int:
    """Dessine une rangee de cartes KPI a partir de (label, valeur, alerte_si_positif).

    Retourne la colonne suivant la derniere carte dessinee (utile pour
    savoir jusqu'ou merger les bannieres au-dessus/en-dessous)."""
    for index, (label, value, is_warning_kpi) in enumerate(kpis):
        col_start = start_col + index * card_width
        col_end = col_start + card_width - 1
        alert = is_warning_kpi and isinstance(value, (int, float)) and value > 0
        fill = WARNING_FILL if alert else ACCENT_LIGHT_FILL
        border = WARNING_BORDER_STYLE if alert else CARD_BORDER_STYLE
        number_font = Font(name=FONT_NAME, bold=True, size=22, color=WARNING_TEXT if alert else NAVY)

        sheet.merge_cells(start_row=row_number, start_column=col_start, end_row=row_number, end_column=col_end)
        sheet.merge_cells(start_row=row_label, start_column=col_start, end_row=row_label, end_column=col_end)

        number_cell = sheet.cell(row=row_number, column=col_start, value=value)
        number_cell.font = number_font
        number_cell.alignment = Alignment(horizontal="center", vertical="bottom")

        label_cell = sheet.cell(row=row_label, column=col_start, value=label)
        label_cell.font = KPI_LABEL_FONT
        label_cell.alignment = Alignment(horizontal="center", vertical="top")

        for row in (row_number, row_label):
            for col in range(col_start, col_end + 1):
                cell = sheet.cell(row=row, column=col)
                cell.fill = fill
                left = border.left if col == col_start else Side(style=None)
                right = border.right if col == col_end else Side(style=None)
                top = border.top if row == row_number else Side(style=None)
                bottom = border.bottom if row == row_label else Side(style=None)
                cell.border = Border(left=left, right=right, top=top, bottom=bottom)

    sheet.row_dimensions[row_number].height = 30
    sheet.row_dimensions[row_label].height = 16
    return start_col + len(kpis) * card_width


# ---------------------------------------------------------------------------
# Synthese - tableau de bord
# ---------------------------------------------------------------------------
def write_summary_sheet(sheet: Worksheet, result: ExtractionResult, page_sheet_names: dict[int, str]) -> None:
    sheet.sheet_view.showGridLines = False

    banner_cols = 10 if result.consolidated_groups else 8

    # Banniere de titre
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=banner_cols)
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=banner_cols)
    banner = sheet.cell(row=1, column=1, value="  Synthese de conversion PDF vers Excel")
    banner.font = TITLE_FONT
    banner.fill = NAVY_FILL
    banner.alignment = Alignment(horizontal="left", vertical="center")
    for row in (1, 2):
        for col in range(1, banner_cols + 1):
            sheet.cell(row=row, column=col).fill = NAVY_FILL
    sheet.row_dimensions[1].height = 22
    sheet.row_dimensions[2].height = 20

    subtitle = sheet.cell(
        row=2,
        column=1,
        value=f"  {result.source_path.name}  ·  genere le {datetime.now().strftime('%d/%m/%Y a %H:%M')}",
    )
    subtitle.font = SUBTITLE_FONT
    subtitle.alignment = Alignment(horizontal="left", vertical="center")

    sheet.row_dimensions[3].height = 8

    # Cartes indicateurs (KPI)
    kpi_row_number, kpi_row_label = 4, 5
    kpis: list[tuple[str, object, bool]] = [
        ("PAGES TRAITEES", len(result.pages), False),
        ("TABLEAUX DETECTES", result.total_tables, False),
        ("LIGNES EXTRAITES", result.total_rows, False),
        ("AVERTISSEMENTS", len(result.warnings), True),
    ]
    if result.consolidated_groups:
        kpis.append(("GROUPES CONSOLIDES", len(result.consolidated_groups), False))
    write_kpi_cards(sheet, kpis, kpi_row_number, kpi_row_label)

    sheet.row_dimensions[6].height = 10

    # Fichier source (discret, sous les cartes)
    source_cell = sheet.cell(row=7, column=1, value=f"Fichier source : {result.source_path}")
    source_cell.font = CAPTION_FONT
    sheet.merge_cells(start_row=7, start_column=1, end_row=7, end_column=banner_cols)
    sheet.row_dimensions[8].height = 6

    if result.consolidated_groups:
        consolidated_note = sheet.cell(
            row=8,
            column=1,
            value="  Voir l'onglet Tableaux_consolides pour les tableaux fusionnes/regroupes par Claude.",
        )
        consolidated_note.font = CAPTION_FONT
        sheet.merge_cells(start_row=8, start_column=1, end_row=8, end_column=banner_cols)
        sheet.row_dimensions[8].height = 14

    # Titre de section + tableau des pages
    section_row = 9
    section_cell = sheet.cell(row=section_row, column=1, value="Detail par page")
    section_cell.font = SECTION_FONT
    sheet.merge_cells(start_row=section_row, start_column=1, end_row=section_row, end_column=banner_cols)
    for col in range(1, banner_cols + 1):
        sheet.cell(row=section_row, column=col).border = ACCENT_BOTTOM_BORDER
    sheet.row_dimensions[section_row].height = 20

    start_row = section_row + 2
    headers = ["Page", "Tableaux", "Lignes", "Statut"]
    for col_index, header in enumerate(headers, start=1):
        cell = sheet.cell(row=start_row, column=col_index, value=header)
        cell.fill = ACCENT_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[start_row].height = 18

    for offset, page in enumerate(result.pages, start=1):
        row_index = start_row + offset
        ok = not page.warnings
        status = "OK" if ok else "A verifier"
        page_cell = sheet.cell(row=row_index, column=1, value=page.page_number)
        page_cell.hyperlink = f"#'{page_sheet_names[page.page_number]}'!A1"
        page_cell.font = HYPERLINK_FONT
        page_cell.alignment = Alignment(horizontal="center")
        sheet.cell(row=row_index, column=2, value=len(page.tables)).alignment = Alignment(horizontal="center")
        sheet.cell(row=row_index, column=3, value=len(page.rows)).alignment = Alignment(horizontal="center")
        status_cell = sheet.cell(row=row_index, column=4, value=f"{status}  ·  {page.extraction_method}")
        status_cell.font = Font(name=FONT_NAME, size=9.5, color=SUCCESS_TEXT if ok else WARNING_TEXT)
        fill = SUCCESS_FILL if ok else WARNING_FILL
        for col_index in (2, 3):
            cell = sheet.cell(row=row_index, column=col_index)
            cell.fill = fill
            cell.font = BODY_FONT
        sheet.cell(row=row_index, column=4).fill = fill

    autosize_columns(sheet, banner_cols)
    sheet.column_dimensions["A"].width = max(sheet.column_dimensions["A"].width or 10, 10)
    sheet.freeze_panes = f"A{start_row + 1}"
    if len(result.pages) > 0:
        sheet.auto_filter.ref = f"A{start_row}:D{start_row + len(result.pages)}"


# ---------------------------------------------------------------------------
# Feuille par page
# ---------------------------------------------------------------------------
def write_page_sheet(
    sheet: Worksheet, page: ExtractedPage, mode: str = "both", used_table_names: set[str] | None = None
) -> None:
    if used_table_names is None:
        used_table_names = set()

    sheet.sheet_view.showGridLines = False

    banner_cols = 8
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=banner_cols - 1)
    title_cell = sheet.cell(row=1, column=1, value=f"  Page {page.page_number}")
    title_cell.font = PAGE_TITLE_FONT
    title_cell.alignment = Alignment(horizontal="left", vertical="center")

    back_cell = sheet.cell(row=1, column=banner_cols, value="Retour a la Synthese")
    back_cell.hyperlink = "#'Synthese'!A1"
    back_cell.font = BACK_LINK_FONT
    back_cell.fill = WHITE_FILL
    back_cell.alignment = Alignment(horizontal="center", vertical="center")
    back_cell.border = CARD_BORDER_STYLE

    for col in range(1, banner_cols):
        sheet.cell(row=1, column=col).fill = ACCENT_FILL
    sheet.row_dimensions[1].height = 24

    stats = f"  {len(page.tables)} tableau(x)  ·  {len(page.rows)} ligne(s)  ·  methode : {page.extraction_method}"
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=banner_cols)
    stats_cell = sheet.cell(row=2, column=1, value=stats)
    stats_cell.font = PAGE_SUBTITLE_FONT
    stats_cell.fill = ACCENT_LIGHT_FILL
    stats_cell.alignment = Alignment(horizontal="left", vertical="center")
    sheet.row_dimensions[2].height = 18

    sheet.row_dimensions[3].height = 8

    current_row = 5
    max_cols = 3
    if mode != "text" and page.tables:
        for index, table in enumerate(page.tables, start=1):
            current_row = write_table_block(sheet, table, index, current_row, used_table_names)
            max_cols = max(max_cols, table.column_count)
            current_row += 2
    elif mode != "text":
        info_cell = sheet.cell(row=current_row, column=1, value="  Aucun tableau structure detecte sur cette page.")
        sheet.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=TEXT_ZONE_SPAN)
        info_cell.fill = WARNING_FILL
        info_cell.font = WARNING_FONT
        info_cell.border = WARNING_BORDER_STYLE
        info_cell.alignment = Alignment(vertical="center")
        sheet.row_dimensions[current_row].height = 20
        current_row += 2

    if mode != "tables":
        current_row = write_text_block(sheet, page, current_row, span_cols=max(max_cols, TEXT_ZONE_SPAN))

    autosize_columns(sheet, max_cols)
    sheet.column_dimensions["A"].width = max(sheet.column_dimensions["A"].width or 10, 20)
    sheet.freeze_panes = "A3"


def write_table_block(
    sheet: Worksheet,
    table: ExtractedTable,
    index: int,
    start_row: int,
    used_table_names: set[str],
    extra_caption: str | None = None,
) -> int:
    title_cell = sheet.cell(row=start_row, column=1, value=f"Tableau {index}")
    title_cell.font = TABLE_TITLE_FONT
    start_row += 1

    caption_text = f"Confiance {table.confidence:.2f}  ·  detection : {detection_method_label(table.detection_method)}"
    if extra_caption:
        caption_text += f"  ·  {extra_caption}"
    caption_cell = sheet.cell(row=start_row, column=1, value=caption_text)
    caption_cell.font = CAPTION_FONT
    start_row += 1

    if table.warnings:
        warning_cell = sheet.cell(row=start_row, column=1, value=" | ".join(table.warnings))
        warning_cell.fill = WARNING_FILL
        warning_cell.font = WARNING_FONT
        warning_cell.border = WARNING_BORDER_STYLE
        start_row += 1

    max_cols = table.column_count
    data_rows = table.rows

    if table.has_header and table.rows:
        headers = make_unique_headers(table.rows[0].values, max_cols)
        data_rows = table.rows[1:]
    else:
        headers = [f"Colonne {i + 1}" for i in range(max_cols)]

    header_row_index = start_row
    for col in range(1, max_cols + 1):
        cell = sheet.cell(row=header_row_index, column=col, value=safe_cell_value(headers[col - 1]))
        cell.fill = ACCENT_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", wrap_text=True, vertical="center")
    sheet.row_dimensions[header_row_index].height = 18

    last_row = header_row_index
    for row_index, row in enumerate(data_rows, start=header_row_index + 1):
        last_row = row_index
        for col_index, value in enumerate(row.values, start=1):
            classified = classify_cell(value)
            cell = sheet.cell(row=row_index, column=col_index, value=classified["value"])
            if classified["kind"] == "text":
                cell.value = safe_cell_value(classified["value"])
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            else:
                cell.number_format = classified["number_format"]
                cell.alignment = Alignment(horizontal="right", vertical="top")
            cell.font = BODY_FONT

    if last_row > header_row_index and max_cols > 0:
        add_excel_table(sheet, table.table_id, header_row_index, last_row, max_cols, used_table_names)

    return start_row + max(len(data_rows) + 1, 1)


def add_excel_table(
    sheet: Worksheet, table_id: str, header_row: int, last_row: int, max_cols: int, used_table_names: set[str]
) -> None:
    """Cree un tableau Excel natif (filtre + bandes de couleur) sur la
    plage donnee, plutot qu'un simple auto_filter qui ne peut exister
    qu'une fois par feuille - insuffisant des qu'une page contient
    plusieurs tableaux."""
    last_col_letter = get_column_letter(max_cols)
    ref = f"A{header_row}:{last_col_letter}{last_row}"

    name = safe_table_name(f"T_{sheet.title}_{table_id}", used_table_names)
    excel_table = Table(displayName=name, ref=ref)
    excel_table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium9",
        showRowStripes=True,
        showFirstColumn=False,
        showLastColumn=False,
        showColumnStripes=False,
    )
    sheet.add_table(excel_table)


def safe_table_name(raw_name: str, used_table_names: set[str]) -> str:
    cleaned = "".join(char if char.isalnum() else "_" for char in raw_name)
    if not cleaned or not (cleaned[0].isalpha() or cleaned[0] == "_"):
        cleaned = f"T_{cleaned}"
    cleaned = cleaned[:60]
    candidate = cleaned
    suffix = 2
    while candidate in used_table_names:
        candidate = f"{cleaned}_{suffix}"
        suffix += 1
    used_table_names.add(candidate)
    return candidate


def detection_method_label(method: str) -> str:
    labels = {
        "lines": "bordures du PDF",
        "regex": "alignement du texte",
        "position": "position des mots",
        "grid": "grille de lignes (image scannee)",
        "claude": "Claude (vision, API)",
        "claude_merge": "Claude (fusion multi-pages)",
    }
    return labels.get(method, method)


# ---------------------------------------------------------------------------
# Feuille "Tableaux_consolides" - resultat de la passe de reorganisation
# document-entier (--smart-merge). Purement additive : n'existe que si des
# groupes ont ete decides, et ne remplace jamais les feuilles par page.
# ---------------------------------------------------------------------------
CONSOLIDATED_SHEET_COLS = 10

MERGED_BADGE_FILL = PatternFill("solid", fgColor=ACCENT)
GROUPED_BADGE_FILL = PatternFill("solid", fgColor="6B8DB8")
BADGE_FONT = Font(name=FONT_NAME, bold=True, size=8.5, color=WHITE)
REASON_FONT = Font(name=FONT_NAME, italic=True, size=9, color=INK)


def write_consolidated_sheet(
    sheet: Worksheet,
    result: ExtractionResult,
    page_sheet_names: dict[int, str],
    used_table_names: set[str],
) -> None:
    sheet.sheet_view.showGridLines = False
    cols = CONSOLIDATED_SHEET_COLS

    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=cols)
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=cols)
    banner = sheet.cell(row=1, column=1, value="  Tableaux consolides (reorganisation Claude)")
    banner.font = TITLE_FONT
    banner.fill = NAVY_FILL
    for col in range(1, cols + 1):
        sheet.cell(row=1, column=col).fill = NAVY_FILL
        sheet.cell(row=2, column=col).fill = NAVY_FILL
    sheet.row_dimensions[1].height = 22
    sheet.row_dimensions[2].height = 18
    subtitle = sheet.cell(
        row=2,
        column=1,
        value="  Tableaux fusionnes ou regroupes par nature, en plus des feuilles par page (rien n'est supprime ni remplace)",
    )
    subtitle.font = SUBTITLE_FONT
    sheet.row_dimensions[3].height = 8

    groups = result.consolidated_groups
    merged_groups = [g for g in groups if g.merged]
    grouped_groups = [g for g in groups if not g.merged]
    tables_involved = sum(len(g.source_table_ids) for g in groups)
    pages_covered = {p for g in groups for p in (_source_pages_from_ids(g.source_table_ids) | {t.page_number for t in g.tables})}

    kpi_row_number, kpi_row_label = 4, 5
    kpis: list[tuple[str, object, bool]] = [
        ("GROUPES", len(groups), False),
        ("TABLEAUX FUSIONNES", len(merged_groups), False),
        ("TABLEAUX REGROUPES", len(grouped_groups), False),
        ("TABLEAUX D'ORIGINE IMPLIQUES", tables_involved, False),
        ("PAGES COUVERTES", len(pages_covered), False),
    ]
    write_kpi_cards(sheet, kpis, kpi_row_number, kpi_row_label, card_width=2)
    sheet.row_dimensions[6].height = 10

    current_row = 8
    for group in groups:
        current_row = write_consolidated_group(sheet, group, current_row, page_sheet_names, used_table_names)
        current_row += 2

    autosize_columns(sheet, cols)
    sheet.column_dimensions["A"].width = max(sheet.column_dimensions["A"].width or 10, 22)
    sheet.freeze_panes = "A7"


def write_consolidated_group(
    sheet: Worksheet,
    group: ConsolidatedGroup,
    start_row: int,
    page_sheet_names: dict[int, str],
    used_table_names: set[str],
) -> int:
    cols = CONSOLIDATED_SHEET_COLS

    # Badge FUSIONNE / REGROUPE - bien visible, separe du titre, pour
    # marquer clairement de quel type de regroupement il s'agit.
    badge_text = "FUSIONNE" if group.merged else "REGROUPE"
    badge_fill = MERGED_BADGE_FILL if group.merged else GROUPED_BADGE_FILL
    badge_cell = sheet.cell(row=start_row, column=1, value=f"  {badge_text}  ")
    badge_cell.font = BADGE_FONT
    badge_cell.fill = badge_fill
    badge_cell.alignment = Alignment(horizontal="center", vertical="center")

    title_cell = sheet.cell(row=start_row, column=2, value=group.title)
    title_cell.font = SECTION_FONT
    sheet.merge_cells(start_row=start_row, start_column=2, end_row=start_row, end_column=cols)
    for col in range(2, cols + 1):
        sheet.cell(row=start_row, column=col).border = ACCENT_BOTTOM_BORDER
    sheet.row_dimensions[start_row].height = 18
    start_row += 1

    source_pages = sorted({table.page_number for table in group.tables} | _source_pages_from_ids(group.source_table_ids))
    pages_label = ", ".join(str(p) for p in source_pages) if source_pages else "?"
    caption_value = f"Categorie : {group.category or 'non precisee'}  ·  pages sources : {pages_label}  ·  {len(group.source_table_ids)} tableau(x) d'origine"
    caption_cell = sheet.cell(row=start_row, column=1, value=caption_value)
    caption_cell.font = CAPTION_FONT
    sheet.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=cols)
    start_row += 1

    if group.reason:
        reason_cell = sheet.cell(row=start_row, column=1, value=f"  → {group.reason}")
        reason_cell.font = REASON_FONT
        sheet.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=cols)
        start_row += 1

    if group.warnings:
        warning_cell = sheet.cell(row=start_row, column=1, value=" | ".join(group.warnings))
        warning_cell.fill = WARNING_FILL
        warning_cell.font = WARNING_FONT
        warning_cell.border = WARNING_BORDER_STYLE
        sheet.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=cols)
        start_row += 1

    start_row += 1
    for index, table in enumerate(group.tables, start=1):
        # Pour un groupe REGROUPE (tableaux distincts), on rappelle la page
        # d'origine de CHAQUE tableau individuellement - pour une fusion, la
        # tracabilite est deja assuree ligne par ligne via la colonne
        # "Page source" ajoutee par merge_tables().
        extra_caption = None if group.merged else f"page source : {table.page_number}"
        start_row = write_table_block(sheet, table, index, start_row, used_table_names, extra_caption=extra_caption)
        start_row += 2

    return start_row


def _source_pages_from_ids(table_ids: list[str]) -> set[int]:
    """Extrait le numero de page depuis un table_id de la forme
    'page_<N>_...' quand la table correspondante n'a pas ete resolue
    (ex: reference introuvable) - purement informatif pour la legende."""
    pages: set[int] = set()
    for table_id in table_ids:
        parts = table_id.split("_")
        if len(parts) >= 2 and parts[0] == "page" and parts[1].isdigit():
            pages.add(int(parts[1]))
    return pages


# ---------------------------------------------------------------------------
# Zone de texte - chaque paragraphe est presente comme une "carte" encadree,
# bien separee des autres, plutot que comme des lignes d'un tableau brut :
# c'est plus lisible pour du texte libre (correspondance, notes, clauses...)
# et evite l'effet grille qui ne convient qu'aux donnees tabulaires.
# ---------------------------------------------------------------------------
def write_text_block(sheet: Worksheet, page: ExtractedPage, start_row: int, span_cols: int = TEXT_ZONE_SPAN) -> int:
    section_cell = sheet.cell(row=start_row, column=1, value="Texte extrait")
    section_cell.font = SECTION_FONT
    sheet.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=span_cols)
    for col in range(1, span_cols + 1):
        sheet.cell(row=start_row, column=col).border = ACCENT_BOTTOM_BORDER
    sheet.row_dimensions[start_row].height = 20
    start_row += 1
    sheet.row_dimensions[start_row].height = 6
    start_row += 1

    text_paragraphs = [row.source_line for row in page.rows if row.kind == "text" and row.source_line.strip()]
    other_rows = [row for row in page.rows if row.kind != "text"]

    if not text_paragraphs and not other_rows:
        empty_cell = sheet.cell(row=start_row, column=1, value="  Aucun texte exploitable detecte sur cette page.")
        sheet.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=span_cols)
        empty_cell.fill = WARNING_FILL
        empty_cell.font = WARNING_FONT
        empty_cell.border = WARNING_BORDER_STYLE
        empty_cell.alignment = Alignment(vertical="center")
        sheet.row_dimensions[start_row].height = 20
        return start_row + 2

    paragraph_number = 0
    for content in text_paragraphs:
        paragraph_number += 1
        label_row = start_row
        label_cell = sheet.cell(row=label_row, column=1, value=f"PARAGRAPHE {paragraph_number}")
        label_cell.font = KPI_LABEL_FONT
        sheet.merge_cells(start_row=label_row, start_column=1, end_row=label_row, end_column=span_cols)
        sheet.row_dimensions[label_row].height = 12

        body_row = label_row + 1
        cleaned = safe_cell_value(content)
        body_cell = sheet.cell(row=body_row, column=1, value=cleaned)
        sheet.merge_cells(start_row=body_row, start_column=1, end_row=body_row, end_column=span_cols)
        body_cell.font = BODY_FONT
        body_cell.alignment = Alignment(wrap_text=True, vertical="top", horizontal="left")
        for col in range(1, span_cols + 1):
            sheet.cell(row=body_row, column=col).border = CARD_BORDER_STYLE
            sheet.cell(row=body_row, column=col).fill = WHITE_FILL

        # Hauteur estimee a partir de la longueur du texte et de la largeur
        # de la carte, pour que le paragraphe reste entierement visible
        # sans avoir a redimensionner la ligne manuellement dans Excel.
        approx_chars_per_line = max(20, span_cols * 16)
        estimated_lines = max(1, -(-len(cleaned) // approx_chars_per_line))
        sheet.row_dimensions[body_row].height = min(15 * estimated_lines + 6, 320)

        start_row = body_row + 1
        sheet.row_dimensions[start_row].height = 8
        start_row += 1

    if other_rows:
        misc_header = sheet.cell(row=start_row, column=1, value="Autres elements")
        misc_header.font = CAPTION_FONT
        sheet.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=span_cols)
        start_row += 1
        for row in other_rows:
            cell = sheet.cell(row=start_row, column=1, value=safe_cell_value(row.source_line))
            sheet.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=span_cols)
            cell.font = BODY_FONT
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            start_row += 1

    return start_row + 1


def write_metadata_sheet(workbook: Workbook, result: ExtractionResult) -> None:
    sheet = workbook.create_sheet(title="Metadonnees")
    sheet.sheet_properties.tabColor = GRAY_TEXT
    sheet.sheet_view.showGridLines = False

    header_cell = sheet.cell(row=1, column=1, value="Metadonnees")
    header_cell.font = SECTION_FONT
    sheet.merge_cells("A1:B1")
    for col in (1, 2):
        sheet.cell(row=1, column=col).border = ACCENT_BOTTOM_BORDER
    sheet.row_dimensions[1].height = 20
    sheet.row_dimensions[2].height = 6

    rows = [
        ("Fichier source", str(result.source_path)),
        ("Date conversion", datetime.now().isoformat(timespec="seconds")),
        ("Version outil", __version__),
        ("Nombre de pages", len(result.pages)),
        ("Tableaux detectes", result.total_tables),
        ("Lignes extraites", result.total_rows),
        ("Avertissements", len(result.warnings)),
    ]
    for offset, (key, value) in enumerate(rows, start=3):
        sheet.cell(row=offset, column=1, value=key).font = BOLD_FONT
        sheet.cell(row=offset, column=2, value=value).font = BODY_FONT
    autosize_columns(sheet, 2)


def write_quality_sheet(workbook: Workbook, result: ExtractionResult, page_sheet_names: dict[int, str]) -> None:
    sheet = workbook.create_sheet(title="Controle_qualite")
    sheet.sheet_view.showGridLines = False

    has_issues = any(page.warnings or any(t.confidence < LOW_CONFIDENCE_THRESHOLD for t in page.tables) for page in result.pages)
    sheet.sheet_properties.tabColor = WARNING_TEXT if has_issues else SUCCESS_TEXT

    headers = ["Page", "Type", "Message"]
    for col_index, header in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=col_index, value=header)
        cell.fill = ACCENT_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 18

    row_index = 2
    for page in result.pages:
        low_confidence_tables = [t for t in page.tables if t.confidence < LOW_CONFIDENCE_THRESHOLD]
        messages = list(page.warnings)
        for table in low_confidence_tables:
            messages.append(f"{table.table_id} : confiance faible ({table.confidence}) - a verifier manuellement.")

        if not messages:
            sheet.cell(row=row_index, column=1, value=page.page_number).font = BODY_FONT
            ok_cell = sheet.cell(row=row_index, column=2, value=f"OK ({page.extraction_method})")
            ok_cell.font = Font(name=FONT_NAME, size=9.5, color=SUCCESS_TEXT)
            sheet.cell(row=row_index, column=3, value="Page traitee").font = BODY_FONT
            for col_index in range(1, 4):
                sheet.cell(row=row_index, column=col_index).fill = SUCCESS_FILL
            row_index += 1
            continue

        for message in messages:
            page_cell = sheet.cell(row=row_index, column=1, value=page.page_number)
            if page.page_number in page_sheet_names:
                page_cell.hyperlink = f"#'{page_sheet_names[page.page_number]}'!A1"
                page_cell.font = HYPERLINK_FONT
            else:
                page_cell.font = BODY_FONT
            type_cell = sheet.cell(row=row_index, column=2, value="A verifier")
            type_cell.font = Font(name=FONT_NAME, size=9.5, color=WARNING_TEXT)
            sheet.cell(row=row_index, column=3, value=message).font = BODY_FONT
            for col_index in range(1, 4):
                sheet.cell(row=row_index, column=col_index).fill = WARNING_FILL
            row_index += 1

    autosize_columns(sheet, 3)
    sheet.freeze_panes = "A2"
    if row_index > 2:
        sheet.auto_filter.ref = f"A1:C{row_index - 1}"


def autosize_columns(sheet: Worksheet, max_cols: int) -> None:
    for col_index in range(1, max_cols + 1):
        letter = get_column_letter(col_index)
        max_length = 10
        for cell in sheet[letter]:
            if cell.value is not None:
                max_length = max(max_length, min(len(str(cell.value)), 60))
        sheet.column_dimensions[letter].width = max_length + 2


def safe_sheet_title(title: str) -> str:
    invalid = ["\\", "/", "*", "?", ":", "[", "]"]
    for char in invalid:
        title = title.replace(char, " ")
    return title[:31]


def make_unique_headers(values: list[str], max_cols: int) -> list[str]:
    headers: list[str] = []
    seen: dict[str, int] = {}
    for index in range(max_cols):
        value = values[index].strip() if index < len(values) and values[index].strip() else f"Colonne {index + 1}"
        count = seen.get(value, 0) + 1
        seen[value] = count
        headers.append(value if count == 1 else f"{value}_{count}")
    return headers


def safe_cell_value(value) -> str:
    if value is None:
        return ""
    text = str(value)
    return "".join(char for char in text if char == "\n" or char == "\t" or ord(char) >= 32)


def save_workbook_safely(workbook: Workbook, output_path: Path) -> Path:
    try:
        workbook.save(output_path)
        return output_path
    except PermissionError:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        fallback_path = output_path.with_name(f"{output_path.stem}_{timestamp}{output_path.suffix}")
        workbook.save(fallback_path)
        return fallback_path
